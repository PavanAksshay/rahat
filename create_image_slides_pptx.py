import os
from pptx import Presentation
from pptx.util import Inches

def build_image_deck(output_path='RAHAT_Inauguration_Presentation.pptx'):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    slide_images = [
        'assets/slides/slide_01_clean.png',
        'assets/slides/slide_02_clean.png',
        'assets/slides/slide_03_clean.png',
        'assets/slides/slide_04_clean.png',
        'assets/slides/slide_05_clean.png'
    ]

    for idx, img_path in enumerate(slide_images, 1):
        if not os.path.exists(img_path):
            fallback = img_path.replace('_clean.png', '.png')
            if os.path.exists(fallback):
                img_path = fallback
            else:
                continue

        slide = prs.slides.add_slide(blank_layout)
        slide.shapes.add_picture(
            img_path,
            left=Inches(0),
            top=Inches(0),
            width=Inches(13.333),
            height=Inches(7.5)
        )
        print(f"Added Slide {idx}: {img_path}")

    prs.save(output_path)
    print(f"Presentation successfully saved to {output_path}")

if __name__ == '__main__':
    build_image_deck('/Users/pavanaksshay/rahat/RAHAT_Inauguration_Presentation.pptx')
    build_image_deck('/Users/pavanaksshay/rahat/RAHAT_Presentation_5_Slides.pptx')
